import io
from docx import Document
from docx.shared import Pt, RGBColor

from app.services.hardening.models import AuditReport
from .base import BaseReportExporter


class DOCXReportExporter(BaseReportExporter):
    def export(self, report: AuditReport) -> bytes:
        doc = Document()

        # Título
        title = doc.add_heading('Reporte Ejecutivo de Hardening', level=0)
        title.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        # Detalles
        p_info = doc.add_paragraph()
        p_info.add_run(f'ID Evaluación: ').bold = True
        p_info.add_run(f'{report.id}\n')
        p_info.add_run(f'Fecha de Ejecución: ').bold = True
        p_info.add_run(f'{report.created_at.strftime("%Y-%m-%d %H:%M:%S")}\n')
        p_info.add_run(f'Puntaje Global: ').bold = True
        p_info.add_run(f'{getattr(report, "score", 0.0):.1f}%')

        doc.add_heading('Detalle de Hallazgos', level=1)

        # Tabla de Hallazgos
        findings = getattr(report, "findings", []) or []
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        headers = ['Regla', 'Estado', 'Severidad', 'Puntaje', 'Razón / Evidencia']
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True

        for f in findings:
            row_cells = table.add_row().cells
            row_cells[0].text = str(f.rule_id)
            row_cells[1].text = f.status.value if hasattr(f.status, 'value') else str(f.status)
            row_cells[2].text = f.severity.value if hasattr(f.severity, 'value') else str(f.severity)
            row_cells[3].text = f"{getattr(f, 'compliance_score', 0.0):.1f}%"
            row_cells[4].text = f.reason or f.raw_evidence or "N/A"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()